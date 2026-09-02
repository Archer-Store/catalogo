import csv
import openpyxl
import requests
from collections import defaultdict
from config import (
    costo_gramo_vellon,
    costo_cm2_microfibra,
    costo_cm2_cordura,
    costo_viaticos,
    PRECIO_HOJA_A3,
    MULTIPLICADOR_COSTO,
    COMISION_ML
)

# --- CONFIGURACIÓN DE NOTION ---
NOTION_TOKEN = "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"
DATABASE_ID = "3c467edf6aa3804bb22eff38cf888fd0"

# 🎯 LISTA DE PRODUCTOS PERMITIDOS
SKUS_PERMITIDOS = {
    'DAKI85', 'FDAKI95D', 'Al2020', 'Al2025', 
    'DAKI6020', 'Al3020', 'FU10S', 'FU10D'
}

# 📐 ESPECIFICACIONES DE MEDIDAS POR SKU (Altura x Ancho)
DETALLES_PRODUCTOS = {
    'DAKI85':   {'altura': 85,  'ancho': 28},
    'FDAKI95D': {'altura': 95,  'ancho': 30},
    'Al2020':   {'altura': 20,  'ancho': 20},
    'Al2025':   {'altura': 20,  'ancho': 25},
    'DAKI6020': {'altura': 60,  'ancho': 20},
    'Al3020':   {'altura': 30,  'ancho': 20},
    'FU10S':    {'altura': 100, 'ancho': 30},
    'FU10D':    {'altura': 100, 'ancho': 30},
}

# ==============================================================================
# CONEXIÓN Y CONSULTA A NOTION
# ==============================================================================

def obtener_variantes_pendientes():
    """Obtiene los diseños pendientes desde Notion."""
    headers = {
        "Authorization": f"Bearer {NOTION_TOKEN}",
        "Content-Type": "application/json",
        "Notion-Version": "2022-06-28"
    }
    url = f"https://api.notion.com/v1/databases/{DATABASE_ID}/query"
    
    payload = {
        "filter": {
            "property": "Estado",
            "status": {"equals": "In progress"}
        }
    }
    
    registros = []
    response = requests.post(url, headers=headers, json=payload)
    
    if response.status_code == 200:
        data = response.json()
        for page in data.get("results", []):
            props = page.get("properties", {})
            page_id = page["id"]
            
            title_list = props.get("Personaje", {}).get("title", [])
            personaje = title_list[0]["text"]["content"] if title_list else None
            
            if not personaje:
                continue

            modelo_list = props.get("Modelo", {}).get("rich_text", [])
            modelo = modelo_list[0]["text"]["content"] if modelo_list else "Modelo 1"

            select_obj = props.get("Franquicia", {}).get("select")
            serie = select_obj["name"] if select_obj else "General"
            
            cover_obj = page.get("cover", {})
            foto_url = ""
            if cover_obj.get("type") == "external":
                foto_url = cover_obj.get("external", {}).get("url", "")
            elif cover_obj.get("type") == "file":
                foto_url = cover_obj.get("file", {}).get("url", "")

            registros.append({
                'id': page_id,
                'personaje': personaje,
                'modelo': modelo,
                'serie': serie,
                'foto': foto_url if foto_url else 'https://i.imgur.com/ejemplo.jpg'
            })
        print(f"📥 Se obtuvieron {len(registros)} variantes pendientes desde Notion.")
    else:
        print(f"⚠️ Error al conectar con Notion ({response.status_code}): {response.text}")

    return registros

def marcar_como_publicado(page_id):
    """Cambia el estado en Notion a 'Publicado MeLi'."""
    headers = {
        "Authorization": f"Bearer {NOTION_TOKEN}",
        "Content-Type": "application/json",
        "Notion-Version": "2022-06-28"
    }
    url = f"https://api.notion.com/v1/pages/{page_id}"
    payload = {
        "properties": {
            "Estado": {"status": {"name": "Publicado MeLi"}}
        }
    }
    requests.patch(url, headers=headers, json=payload)

# ==============================================================================
# CÁLCULOS Y FORMATO
# ==============================================================================

def calcular_precio_ml(precio_sugerido):
    COMISION = 0.266
    precio_base = (precio_sugerido + 2850) / (1 - COMISION)
    if precio_base >= 33000:
        precio_final = (precio_sugerido + 8250) / (1 - COMISION)
    else:
        precio_final = precio_base
    return round(precio_final)

def armar_titulo_seo(tipo_producto, personaje, serie):
    titulo = f"{tipo_producto} {personaje} {serie}".strip()
    if len(titulo) <= 60:
        return titulo
    titulo = f"{tipo_producto} {personaje}".strip()
    if len(titulo) <= 60:
        return titulo
    return titulo[:60].strip()

def cargar_csv_autodetect(ruta):
    with open(ruta, mode='r', encoding='utf-8-sig') as f:
        primera_linea = f.readline()
        delim = ';' if ';' in primera_linea else ','
        f.seek(0)
        lector = csv.DictReader(f, delimiter=delim)
        return [{k.strip(): (v.strip() if v else v) for k, v in fila.items() if k} for fila in lector]

# ==============================================================================
# GENERACIÓN DE CATÁLOGO CON VARIANTES
# ==============================================================================

def generar_catalogo_masivo():
    productos = cargar_csv_autodetect('productos.csv')
    registros = obtener_variantes_pendientes()
    
    if not registros:
        print("ℹ️ No hay variaciones pendientes en Notion.")
        return [], []

    # Agrupar por (personaje, serie)
    agrupados = defaultdict(list)
    for r in registros:
        agrupados[(r['personaje'], r['serie'])].append(r)

    publicaciones = []

    for prod in productos:
        codigo_sku = prod.get('codigo')
        if codigo_sku not in SKUS_PERMITIDOS:
            continue

        costo_insumo = (float(prod['cant_vellon_g']) * costo_gramo_vellon) + float(prod['costo_insumo_directo'])
        costo_tela = (float(prod['cm2_microfibra']) * costo_cm2_microfibra) + (float(prod['cm2_cordura']) * costo_cm2_cordura)
        costo_hojas = float(prod['cant_hojas']) * PRECIO_HOJA_A3
        costo_total = costo_insumo + costo_tela + costo_hojas
        
        if costo_total == 0:
            continue
            
        precio_sugerido = (costo_total * MULTIPLICADOR_COSTO) + costo_viaticos
        precio_ml = calcular_precio_ml(precio_sugerido)
        
        detalles_sku = DETALLES_PRODUCTOS.get(codigo_sku, {})
        ancho_final = detalles_sku.get('ancho', prod.get('largo', 30))
        altura_final = detalles_sku.get('altura', prod.get('altura', 100))

        for (personaje, serie), lista_variantes in agrupados.items():
            titulo = armar_titulo_seo(prod['producto'], personaje, serie)
            tag_personaje = personaje.replace(' ', '').upper()[:5]

            # 📝 PLANTILLA DE DESCRIPCIÓN PERSONALIZADA
            descripcion = (
                f'Almohada DakImakura "{personaje}" de "{serie}" Personalizada (Imagen de un lado) de {altura_final}cm x {ancho_final}cm\n\n'
                f'Incluye relleno (no es desmontable)\n'
                f'¡Descubrí la máxima comodidad y estilo con nuestras dakimakuras de anime!\n\n'
                f'Elegí tu diseño, mandanos mensaje cuando hayas realizado tu compra y nosotros hacemos realidad tu personaje!\n\n'
                f' Transforma tu espacio en un rincón lleno de magia\n'
                f' Sumergite en abrazos cómodos mientras tu personaje favorito te acompaña en tus sueños.\n'
                f'Disponibles en diseños cautivadores personalizados y materiales de calidad, estas dakimakuras son el complemento perfecto para cualquier amante del anime.\n\n'
                f' ¡Pedinos el diseño de tu personaje favorito y que tus noches sean épicas con nuestra colección, solo para verdaderos fanáticos!\n\n'
                f'ARCHER.ESTAMPADOS'
            )

            for idx, var in enumerate(lista_variantes, start=1):
                sku_variante = f"{codigo_sku}-{tag_personaje}-V{idx}"

                publicaciones.append({
                    'sku': sku_variante,
                    'titulo': titulo,
                    'personaje': personaje,
                    'serie': serie,
                    'precio_lista': round(precio_sugerido),
                    'precio_ml': precio_ml,
                    'largo_titulo': len(titulo),
                    'color': 'Blanco',
                    'fotos': var['foto'],
                    'descripcion': descripcion,
                    'marca': prod.get('marca', 'Genérica'),
                    'modelo': var['modelo'],
                    'largo': ancho_final,
                    'altura': altura_final,
                    'relleno': prod.get('relleno', 'Fibra siliconada')
                })
            
    return publicaciones, registros

# ==============================================================================
# EXPORTACIÓN DE ARCHIVOS
# ==============================================================================

def exportar_a_excel_ml(publicaciones, plantilla_path='Publicar-08-05-12_51_40.xlsx', ruta_salida='publicaciones_ml_listo.xlsx'):
    if not publicaciones:
        print("⚠️ No hay publicaciones para exportar a Excel.")
        return False

    try:
        wb = openpyxl.load_workbook(plantilla_path)
        ws = wb["Almohadas"] if "Almohadas" in wb.sheetnames else wb.active
    except FileNotFoundError:
        print(f"⚠️ No se encontró la plantilla '{plantilla_path}'.")
        return False

    fila_inicial = 9

    for idx, p in enumerate(publicaciones):
        fila = fila_inicial + idx

        ws.cell(row=fila, column=2, value=p['titulo'])
        ws.cell(row=fila, column=3, value=p['largo_titulo'])
        ws.cell(row=fila, column=4, value="Nuevo")
        ws.cell(row=fila, column=5, value="El producto no tiene código registrado")
        ws.cell(row=fila, column=6, value=p['color'])
        ws.cell(row=fila, column=7, value=p['fotos'])
        ws.cell(row=fila, column=8, value=p['sku'])
        ws.cell(row=fila, column=9, value=10)
        ws.cell(row=fila, column=10, value=p['precio_ml'])
        ws.cell(row=fila, column=11, value="Unidad")
        ws.cell(row=fila, column=12, value=1)
        ws.cell(row=fila, column=13, value=p['descripcion'])
        ws.cell(row=fila, column=15, value="Agregar cuotas")
        ws.cell(row=fila, column=17, value="Mercado Envíos")
        ws.cell(row=fila, column=18, value="A cargo del comprador")
        ws.cell(row=fila, column=19, value="Acepto")
        ws.cell(row=fila, column=20, value="Garantía del vendedor")
        ws.cell(row=fila, column=21, value=30)
        ws.cell(row=fila, column=22, value="días")
        ws.cell(row=fila, column=23, value="No ofrezco")
        ws.cell(row=fila, column=24, value=p['marca'])
        ws.cell(row=fila, column=25, value=p['modelo'])
        ws.cell(row=fila, column=26, value=p['largo'])
        ws.cell(row=fila, column=27, value="cm")
        ws.cell(row=fila, column=28, value=p['altura'])
        ws.cell(row=fila, column=29, value="cm")
        ws.cell(row=fila, column=30, value=p['relleno'])
        ws.cell(row=fila, column=31, value="No")
        ws.cell(row=fila, column=32, value="Tradicional")
        ws.cell(row=fila, column=33, value=15)
        ws.cell(row=fila, column=34, value="cm")
        ws.cell(row=fila, column=35, value="No")
        ws.cell(row=fila, column=36, value="Sí")
        ws.cell(row=fila, column=37, value="Sí")
        ws.cell(row=fila, column=38, value="Sí")
        ws.cell(row=fila, column=39, value="No")
        ws.cell(row=fila, column=40, value="Sí")

    wb.save(ruta_salida)
    print(f"📊 ¡Excel listo guardado con variantes en: {ruta_salida}!")
    return True

if __name__ == "__main__":
    resultado, lista_registros = generar_catalogo_masivo()
    
    if resultado:
        print("=" * 95)
        print(f"🚀 SE GENERARON {len(resultado)} FILAS DE VARIACIONES PARA MERCADOLIBRE")
        print("=" * 95)
        
        exito = exportar_a_excel_ml(resultado, plantilla_path='Publicar-08-05-12_51_40.xlsx')
        
        if exito:
            print("🔄 Actualizando estado en Notion...")
            for reg in lista_registros:
                marcar_como_publicado(reg['id'])
            print("✅ Notion actualizado.")